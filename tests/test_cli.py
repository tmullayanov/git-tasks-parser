import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from main import DATE_FORMAT, ExportError, Repository, export


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.repo = self.root / "repo"
        self.repo.mkdir()
        self.env = os.environ.copy()
        self.env.update({
            "GIT_AUTHOR_NAME": "Author",
            "GIT_AUTHOR_EMAIL": "=1+1",
            "GIT_AUTHOR_DATE": "2020-01-01T00:00:00+00:00",
            "GIT_COMMITTER_NAME": "Committer",
            "GIT_COMMITTER_EMAIL": "committer@example.com",
            "GIT_COMMITTER_DATE": "2026-10-28T11:30:45+05:00",
        })
        self.git("init", "-b", "main")

    def git(self, *args):
        result = subprocess.run(
            ["git", "-C", str(self.repo), *args],
            env=self.env, capture_output=True, text=True, check=True,
        )
        return result.stdout.strip()

    def commit(self, message, body=None):
        args = ["commit", "--allow-empty", "-m", message]
        if body is not None:
            args.extend(["-m", body])
        return self.git(*args)

    def test_merge_keys_author_and_text_cells(self):
        self.commit("No issue here")
        root_hash = self.git("rev-parse", "HEAD")
        self.git("switch", "-c", "feature")
        self.commit("Feature", "PROJ-123 OPS-7 PROJ-123 proj-2 xPROJ-6 PROJ-0")
        feature_hash = self.git("rev-parse", "HEAD")
        self.git("switch", "main")
        self.commit("MAIN_2-4")
        main_hash = self.git("rev-parse", "HEAD")
        self.git("merge", "--no-ff", "feature", "-m", "Merge feature")
        merge_hash = self.git("rev-parse", "HEAD")
        self.git("switch", "-c", "unmerged")
        self.commit("HIDDEN-1")
        self.git("switch", "main")

        output = self.root / "report.xlsx"
        self.assertEqual(export([Repository(self.repo)], "main", "Europe/Moscow", output), 5)
        workbook = load_workbook(output)
        self.assertEqual(workbook.sheetnames, ["Commits"])
        rows = list(workbook.active.values)
        self.assertEqual(rows[0], ("Commit hash", "Issue key", "author email", "Commit date", "repository name"))
        self.assertEqual(rows[1], (merge_hash, None, "=1+1", datetime(2026, 10, 28, 9, 30, 45), "repo"))
        self.assertEqual(
            {(hash_, key) for hash_, key, _, _, _ in rows[1:]},
            {
                (merge_hash, None), (main_hash, "MAIN_2-4"),
                (feature_hash, "PROJ-123"), (feature_hash, "OPS-7"),
                (root_hash, None),
            },
        )
        self.assertEqual(rows[-1][0], root_hash)
        self.assertTrue(all(email == "=1+1" for _, _, email, _, _ in rows[1:]))
        self.assertEqual(workbook.active["C2"].data_type, "s")
        self.assertEqual(workbook.active["D2"].data_type, "d")
        self.assertEqual(workbook.active["D2"].number_format, DATE_FORMAT)
        self.assertEqual(workbook.active["E2"].data_type, "s")

        with self.assertRaisesRegex(ExportError, "already exists"):
            export([Repository(self.repo)], "main", "Europe/Moscow", output)
        self.assertEqual(list(load_workbook(output).active.values), rows)

    def test_invalid_branch_and_detached_head(self):
        self.commit("ROOT-1")
        self.git("update-ref", "refs/remotes/origin/main", "HEAD")
        self.git("checkout", "--detach")
        output = self.root / "report.xlsx"
        self.assertEqual(export([Repository(self.repo)], "main", "Europe/Moscow", output), 1)
        self.assertEqual(
            export([Repository(self.repo)], "refs/remotes/origin/main", "Europe/Moscow", self.root / "remote.xlsx"), 1
        )
        with self.assertRaises(ExportError):
            export([Repository(self.repo)], "HEAD", "Europe/Moscow", self.root / "head.xlsx")
        with self.assertRaises(ExportError):
            export([Repository(self.repo)], "refs/tags/v1", "Europe/Moscow", self.root / "tag.xlsx")
        self.assertFalse((self.root / "head.xlsx").exists())

    def test_multiple_repositories_and_time_zones(self):
        self.commit("PROJ-123")
        root_hash = self.git("rev-parse", "HEAD")
        self.commit("OPS-7")
        other = self.root / "other"
        subprocess.run(["git", "clone", "--quiet", str(self.repo), str(other)], check=True)
        tip_hash = self.git("rev-parse", "HEAD")
        repositories = [Repository(self.repo, "backend"), Repository(other, "frontend")]

        output = self.root / "combined.xlsx"
        self.assertEqual(export(repositories, "main", "Asia/Vladivostok", output), 4)
        rows = list(load_workbook(output).active.values)
        self.assertEqual([row[0] for row in rows[1:]], [tip_hash, root_hash, tip_hash, root_hash])
        self.assertEqual([row[4] for row in rows[1:]], ["backend", "backend", "frontend", "frontend"])
        self.assertEqual([row[3] for row in rows[1:]], [datetime(2026, 10, 28, 16, 30, 45)] * 4)

        with self.assertRaisesRegex(ExportError, "Duplicate repository name"):
            export([Repository(self.repo, "Backend"), Repository(other, "backend")], "main", "UTC", self.root / "duplicate.xlsx")
        self.assertFalse((self.root / "duplicate.xlsx").exists())
        with self.assertRaisesRegex(ExportError, "Unknown IANA time zone"):
            export(repositories, "main", "No/Such_Zone", self.root / "invalid-zone.xlsx")

        subprocess.run(["git", "-C", str(other), "branch", "-m", "main", "other-main"], check=True)
        with self.assertRaisesRegex(ExportError, "Branch does not exist"):
            export(repositories, "main", "UTC", self.root / "missing-branch.xlsx")
        self.assertFalse((self.root / "missing-branch.xlsx").exists())

    def test_cli_repository_name_must_follow_repo(self):
        self.commit("PROJ-123")
        other = self.root / "other"
        subprocess.run(["git", "clone", "--quiet", str(self.repo), str(other)], check=True)
        output = self.root / "cli.xlsx"
        command = [
            sys.executable, str(Path(__file__).resolve().parents[1] / "main.py"),
            "--repo", str(self.repo), "--main-branch", "main",
            "--repository-name", "backend", "--timezone", "UTC",
            "--output", str(output),
        ]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("immediately follow", result.stderr)
        self.assertFalse(output.exists())

        command = [
            sys.executable, str(Path(__file__).resolve().parents[1] / "main.py"),
            "--repo", str(self.repo), "--repository-name", "backend",
            "--repo", str(other), "--main-branch", "main",
            "--timezone", "Europe/Moscow", "--output", str(output),
        ]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual([row[4] for row in list(load_workbook(output).active.values)[1:]], ["backend", "other"])


if __name__ == "__main__":
    unittest.main()
